"""Excel de asistencia: una hoja por curso, con formato listo para imprimir o compartir."""
import base64
import io
import re
from datetime import date
from functools import lru_cache

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

AZUL, AZUL_CLARO = "1F4EB4", "E8EFFC"
FILL = {
    "asistencia": ("E5F6EC", "0B7A43", "✓"),
    "falta": ("FDE8E6", "B4231A", "F"),
    "retardo": ("FFF3D1", "8A5A00", "R"),
    "justificada": ("E6F0FD", "1559AD", "J"),
}
DIAS = ["L", "M", "X", "J", "V", "S", "D"]
DIAS_TXT = {"L": "Lun", "M": "Mar", "X": "Mié", "J": "Jue", "V": "Vie", "S": "Sáb", "D": "Dom"}
LADO = Side(style="thin", color="D3DBE8")
BORDE = Border(left=LADO, right=LADO, top=LADO, bottom=LADO)


_font = lru_cache(maxsize=None)(Font)
_alin = lru_cache(maxsize=None)(Alignment)


@lru_cache(maxsize=None)
def _f(hexcolor: str) -> PatternFill:
    return PatternFill("solid", start_color=hexcolor, end_color=hexcolor)


def _nombre_hoja(nombre: str, usados: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "-", nombre).strip()[:28] or "Curso"
    n, i = base, 2
    while n.lower() in usados:
        n = f"{base[:25]} {i}"
        i += 1
    usados.add(n.lower())
    return n


def logo_bytes(data_url: str | None) -> bytes | None:
    """El logo se guarda como data URL (PNG/JPG) en la configuración."""
    if not data_url or "," not in data_url:
        return None
    try:
        return base64.b64decode(data_url.split(",", 1)[1])
    except Exception:  # noqa: BLE001 - un logo dañado no debe impedir el Excel
        return None


def _poner_logo(ws, logo: bytes | None, ancho_px: int, alto_px: int) -> bool:
    """Coloca el logo en A1 ajustado a la caja ancho_px × alto_px. Devuelve True si se puso."""
    if not logo:
        return False
    try:
        img = XLImage(io.BytesIO(logo))
        k = min(ancho_px / img.width, alto_px / img.height)
        img.width, img.height = int(img.width * k), int(img.height * k)
        ws.add_image(img, "A1")
        return True
    except Exception:  # noqa: BLE001
        return False


def _color_pct(p):
    if p is None:
        return ("F2F4F7", "667085")
    if p >= 90:
        return ("E5F6EC", "0B7A43")
    if p >= 75:
        return ("FFF3D1", "8A5A00")
    return ("FDE8E6", "B4231A")


def _pct(c: dict) -> float | None:
    total = c["asistencia"] + c["falta"] + c["retardo"] + c["justificada"]
    return round((c["asistencia"] + c["retardo"] + c["justificada"]) * 100 / total, 1) if total else None


def hoja_curso(wb: Workbook, nombre_hoja: str, escuela: str, curso: dict, alumnos: list[dict],
               fechas: list[date], rango: str, logo: bytes | None = None) -> float | None:
    """alumnos: [{matricula, nombre, retirado, reg: {fecha: estado}}]. Devuelve el % promedio del curso."""
    ws = wb.create_sheet(nombre_hoja)
    ws.sheet_view.showGridLines = False
    ncol_f = len(fechas)
    ult = 3 + max(ncol_f, 1) + 5  # # + matrícula + alumno + fechas + 5 totales
    ult = max(ult, 9)

    con_logo = _poner_logo(ws, logo, 118, 50)
    ini = 3 if con_logo else 1
    if con_logo:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        ws.cell(row=1, column=1).fill = _f("FFFFFF")
    ws.merge_cells(start_row=1, start_column=ini, end_row=1, end_column=ult)
    c = ws.cell(row=1, column=ini, value=escuela)
    c.font, c.fill = _font(bold=True, size=15, color="FFFFFF"), _f(AZUL)
    c.alignment = _alin(vertical="center", indent=1)
    ws.row_dimensions[1].height = 42 if con_logo else 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ult)
    c = ws.cell(row=2, column=1, value=f"Lista de asistencia · {curso['nombre']} ({curso['codigo']})")
    c.font = _font(bold=True, size=13, color="162033")
    c.alignment = _alin(vertical="center", indent=1)
    ws.row_dimensions[2].height = 24

    dias = ", ".join(DIAS_TXT.get(d, d) for d in (curso.get("dias_clase") or "").split(",") if d) or "—"
    info = [("Instructor", curso.get("instructor") or "—"), ("Plantel", curso.get("plantel") or "—"),
            ("Horario", f"{dias} · {curso.get('horario') or '—'}"), ("Periodo", rango)]
    col = 1
    for etiqueta, valor in info:
        ws.cell(row=3, column=col, value=etiqueta.upper()).font = _font(size=8, bold=True, color="667085")
        ws.cell(row=4, column=col, value=valor).font = _font(size=10, bold=True, color="162033")
        col += 3 if col == 1 else 3
    ws.row_dimensions[4].height = 20

    cab = 6
    heads = ["#", "Matrícula", "Alumno"] + [f"{f.day:02d}/{f.month:02d}\n{DIAS_TXT[DIAS[f.weekday()]]}" for f in fechas] \
        + ["Asist.", "Faltas", "Retardos", "Justif.", "% Asist."]
    for j, h in enumerate(heads, 1):
        x = ws.cell(row=cab, column=j, value=h)
        x.font, x.fill, x.border = _font(bold=True, color="FFFFFF", size=10), _f(AZUL), BORDE
        x.alignment = _alin(horizontal="center" if j != 3 else "left", vertical="center", wrap_text=True)
    ws.row_dimensions[cab].height = 34

    pcts, r = [], cab + 1
    por_fecha = {f: 0 for f in fechas}
    for n, a in enumerate(alumnos, 1):
        cnt = {k: 0 for k in FILL}
        zebra = _f("F7F9FC") if n % 2 == 0 else None
        vals = [n, a["matricula"], a["nombre"] + (" (retirado)" if a.get("retirado") else "")]
        for j, v in enumerate(vals, 1):
            x = ws.cell(row=r, column=j, value=v)
            x.border, x.font = BORDE, _font(size=10, color="667085" if a.get("retirado") else "162033", bold=(j == 3))
            x.alignment = _alin(horizontal="left" if j == 3 else "center", vertical="center")
            if j == 2:
                x.font = _font(name="Consolas", size=10)
            if zebra:
                x.fill = zebra
        for k, f in enumerate(fechas):
            est = a["reg"].get(f)
            x = ws.cell(row=r, column=4 + k)
            x.border, x.alignment = BORDE, _alin(horizontal="center", vertical="center")
            if est:
                bg, fg, sym = FILL[est]
                x.value, x.fill, x.font = sym, _f(bg), _font(bold=True, color=fg, size=11)
                cnt[est] += 1
                if est != "falta":
                    por_fecha[f] += 1
            elif zebra:
                x.fill = zebra
        p = _pct(cnt)
        if p is not None:
            pcts.append(p)
        base = 4 + ncol_f
        for k, v in enumerate([cnt["asistencia"], cnt["falta"], cnt["retardo"], cnt["justificada"]]):
            x = ws.cell(row=r, column=base + k, value=v)
            x.border, x.alignment, x.font = BORDE, _alin(horizontal="center"), _font(size=10)
        x = ws.cell(row=r, column=base + 4, value=(p / 100 if p is not None else "—"))
        bg, fg = _color_pct(p)
        x.fill, x.font, x.border = _f(bg), _font(bold=True, color=fg, size=10), BORDE
        x.alignment = _alin(horizontal="center")
        if p is not None:
            x.number_format = "0.0%"
        ws.row_dimensions[r].height = 21
        r += 1

    if not fechas:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ult)
        x = ws.cell(row=r, column=1, value="Todavía no hay asistencias registradas en este periodo.")
        x.font, x.alignment = _font(italic=True, color="667085"), _alin(horizontal="center")
        r += 1
    else:
        ws.cell(row=r, column=3, value="Presentes por día").font = _font(bold=True, size=10)
        ws.cell(row=r, column=3).alignment = _alin(horizontal="right")
        for k, f in enumerate(fechas):
            x = ws.cell(row=r, column=4 + k, value=por_fecha[f])
            x.font, x.fill, x.border = _font(bold=True, size=10, color=AZUL), _f(AZUL_CLARO), BORDE
            x.alignment = _alin(horizontal="center")
        prom = round(sum(pcts) / len(pcts), 1) if pcts else None
        ws.cell(row=r, column=4 + ncol_f + 3, value="Promedio").font = _font(bold=True, size=10)
        ws.cell(row=r, column=4 + ncol_f + 3).alignment = _alin(horizontal="right")
        x = ws.cell(row=r, column=4 + ncol_f + 4, value=(prom / 100 if prom is not None else "—"))
        bg, fg = _color_pct(prom)
        x.fill, x.font, x.border = _f(bg), _font(bold=True, color=fg, size=10), BORDE
        x.alignment = _alin(horizontal="center")
        if prom is not None:
            x.number_format = "0.0%"
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Leyenda:").font = _font(bold=True, size=9, color="667085")
    pos = 3
    for est, txt in (("asistencia", "✓ Asistió"), ("falta", "F Faltó"), ("retardo", "R Retardo"), ("justificada", "J Justif.")):
        bg, fg, _ = FILL[est]
        x = ws.cell(row=r, column=pos, value=txt)
        x.fill, x.font, x.alignment = _f(bg), _font(bold=True, size=9, color=fg), _alin(horizontal="center")
        pos += 1 if pos > 3 else 1
    ws.cell(row=r + 1, column=1, value="El porcentaje cuenta como presencia: asistencias, retardos y faltas justificadas.").font = _font(size=8, italic=True, color="667085")

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 13
    ws.column_dimensions["C"].width = 34
    for k in range(ncol_f):
        ws.column_dimensions[get_column_letter(4 + k)].width = 7.5
    for k in range(5):
        ws.column_dimensions[get_column_letter(4 + ncol_f + k)].width = 10
    ws.freeze_panes = ws.cell(row=cab + 1, column=4)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"{cab}:{cab}"
    return round(sum(pcts) / len(pcts), 1) if pcts else None


def libro(escuela: str, cursos: list[dict], rango: str, logo: bytes | None = None) -> bytes:
    """cursos: [{curso: dict, alumnos: [...], fechas: [date]}]"""
    wb = Workbook()
    resumen = wb.active
    resumen.title = "Resumen"
    usados = {"resumen"}
    filas = []
    for c in cursos:
        nombre = _nombre_hoja(c["curso"]["codigo"] or c["curso"]["nombre"], usados)
        prom = hoja_curso(wb, nombre, escuela, c["curso"], c["alumnos"], c["fechas"], rango, logo)
        filas.append((nombre, c, prom))

    if len(cursos) == 1:
        wb.remove(resumen)
        return _bytes(wb)

    resumen.sheet_view.showGridLines = False
    con_logo = _poner_logo(resumen, logo, 215, 50)
    ini = 2 if con_logo else 1
    if con_logo:
        resumen["A1"].fill = _f("FFFFFF")
    resumen.merge_cells(start_row=1, start_column=ini, end_row=1, end_column=7)
    x = resumen.cell(row=1, column=ini)
    x.value, x.font, x.fill = escuela, _font(bold=True, size=15, color="FFFFFF"), _f(AZUL)
    x.alignment = _alin(vertical="center", indent=1)
    resumen.row_dimensions[1].height = 42 if con_logo else 30
    resumen.merge_cells("A2:G2")
    resumen["A2"].value = f"Resumen de asistencia por curso · {rango}"
    resumen["A2"].font = _font(bold=True, size=13)
    resumen.row_dimensions[2].height = 24
    for j, h in enumerate(["Curso", "Código", "Instructor", "Plantel", "Alumnos", "Clases registradas", "% Asistencia"], 1):
        x = resumen.cell(row=4, column=j, value=h)
        x.font, x.fill, x.border = _font(bold=True, color="FFFFFF"), _f(AZUL), BORDE
        x.alignment = _alin(horizontal="center", vertical="center", wrap_text=True)
    resumen.row_dimensions[4].height = 30
    for i, (hoja, c, prom) in enumerate(filas, 5):
        cu = c["curso"]
        vals = [cu["nombre"], cu["codigo"], cu.get("instructor") or "—", cu.get("plantel") or "—",
                len(c["alumnos"]), len(c["fechas"]), (prom / 100 if prom is not None else "—")]
        for j, v in enumerate(vals, 1):
            x = resumen.cell(row=i, column=j, value=v)
            x.border, x.alignment = BORDE, _alin(horizontal="left" if j <= 4 else "center", vertical="center")
            x.font = _font(size=10)
            if i % 2 == 0:
                x.fill = _f("F7F9FC")
        a = resumen.cell(row=i, column=1)
        a.hyperlink = f"#'{hoja}'!A1"
        a.font = _font(size=10, bold=True, color=AZUL, underline="single")
        p = resumen.cell(row=i, column=7)
        bg, fg = _color_pct(prom)
        p.fill, p.font = _f(bg), _font(bold=True, color=fg)
        if prom is not None:
            p.number_format = "0.0%"
        resumen.row_dimensions[i].height = 22
    for col, w in zip("ABCDEFG", (34, 12, 26, 20, 10, 12, 13)):
        resumen.column_dimensions[col].width = w
    resumen.freeze_panes = "A5"
    resumen.page_setup.orientation = "landscape"
    resumen.page_setup.fitToWidth, resumen.page_setup.fitToHeight = 1, 0
    resumen.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    return _bytes(wb)


def _bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
