"""Excel de reportes con el mismo diseño que la asistencia: logo, encabezado azul, filas alternadas,
formatos de dinero/fecha/porcentaje, estados en color y listo para imprimir."""
import io
import re
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from .asistencia_excel import AZUL, AZUL_CLARO, BORDE, _alin, _f, _font, _poner_logo
from .utils import ahora

_COLOR_METODO = {
    "efectivo": ("E5F6EC", "0B7A43"), "transferencia": ("E6F0FD", "1559AD"), "tarjeta": ("F1E8FD", "6B2FC4"),
    "otro": ("EEF0F4", "475467"),
}
ROJO = ("FDE8E6", "B4231A")
AMARILLO = ("FFF3D1", "8A5A00")
VERDE = ("E5F6EC", "0B7A43")


def _color_estado(v) -> tuple[str, str] | None:
    t = str(v or "").lower()
    if re.search(r"vencid|baja|cancel|falt", t):
        return ROJO
    if re.search(r"pendiente|parcial|por vencer|próximo|proximo|retardo", t):
        return AMARILLO
    if re.search(r"pagad|aplicad|activ|corriente|terminad|asistencia", t):
        return VERDE
    return None


def reporte_excel(rep: dict, escuela: str, logo: bytes | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\\/*?:\[\]]", "-", rep["titulo"])[:30]
    ws.sheet_view.showGridLines = False
    cols = rep["columnas"]
    n = len(cols)
    ult = max(n, 4)

    # ---- banner con logo
    con_logo = _poner_logo(ws, logo, 118, 50)
    ini = 2 if con_logo else 1
    if con_logo:
        ws["A1"].fill = _f("FFFFFF")
    ws.merge_cells(start_row=1, start_column=ini, end_row=1, end_column=ult)
    x = ws.cell(row=1, column=ini, value=escuela)
    x.font, x.fill, x.alignment = _font(bold=True, size=15, color="FFFFFF"), _f(AZUL), _alin(vertical="center", indent=1)
    ws.row_dimensions[1].height = 42 if con_logo else 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ult)
    x = ws.cell(row=2, column=1, value=rep["titulo"])
    x.font, x.alignment = _font(bold=True, size=14, color="162033"), _alin(vertical="center", indent=1)
    ws.row_dimensions[2].height = 26

    sub = rep.get("subtitulo") or ""
    gen = f"Generado el {ahora().strftime('%d/%m/%Y %H:%M')}"
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ult)
    x = ws.cell(row=3, column=1, value=f"{sub}  ·  {gen}" if sub else gen)
    x.font, x.alignment = _font(size=10, color="667085"), _alin(indent=1)

    cab = 5

    def encabezado(r, columnas):
        for j, c in enumerate(columnas, 1):
            x = ws.cell(row=r, column=j, value=c["label"])
            x.font, x.fill, x.border = _font(bold=True, color="FFFFFF", size=10), _f(AZUL), BORDE
            x.alignment = _alin(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 28

    def filas_(r, columnas, filas):
        """Escribe las filas desde r+1 y devuelve la última fila usada."""
        for i, fila in enumerate(filas):
            r += 1
            zebra = _f("F7F9FC") if i % 2 else None
            for j, c in enumerate(columnas, 1):
                v = fila.get(c["key"])
                tipo = c.get("tipo")
                if tipo == "pct" and isinstance(v, (int, float)):
                    v = v / 100
                x = ws.cell(row=r, column=j, value=v)
                x.border, x.font = BORDE, _font(size=10)
                h = "left"
                if tipo == "money":
                    x.number_format, h = '"$"#,##0.00', "right"
                elif tipo == "date" or isinstance(v, (date, datetime)):
                    x.number_format, h = "DD/MM/YYYY", "center"
                elif tipo == "pct":
                    x.number_format, h = "0.0%", "center"
                elif isinstance(v, (int, float)):
                    h = "center"
                if v is None and tipo == "pct":
                    x.value, h = "—", "center"
                x.alignment = _alin(horizontal=h, vertical="center", wrap_text=(h == "left"))
                if zebra:
                    x.fill = zebra
                if "estado" in c["key"]:
                    col = _color_estado(v)
                    if col:
                        x.fill, x.font, x.alignment = _f(col[0]), _font(bold=True, size=10, color=col[1]), _alin(horizontal="center", vertical="center")
                elif tipo == "pct" and isinstance(v, (int, float)):
                    col = VERDE if v >= 0.9 else AMARILLO if v >= 0.75 else ROJO
                    x.fill, x.font = _f(col[0]), _font(bold=True, size=10, color=col[1])
            ws.row_dimensions[r].height = 20
        return r

    por_metodo = any(c["key"] == "metodo" for c in cols) and any(c["key"] == "importe" for c in cols) and rep["filas"]
    r = cab
    resumen_metodos = []
    if por_metodo:
        # ---- una tabla por método de pago
        cols_t = [c for c in cols if c["key"] != "metodo"]
        i_imp = next(j for j, c in enumerate(cols_t, 1) if c["key"] == "importe")
        grupos: dict[str, list] = {}
        for f in rep["filas"]:
            grupos.setdefault(str(f.get("metodo") or "Sin método"), []).append(f)
        orden = sorted(grupos, key=lambda m: (m.lower() != "efectivo", m.lower() == "otro", m.lower()))
        r = cab - 1
        for metodo in orden:
            filas = grupos[metodo]
            validas = [f for f in filas if "cancel" not in str(f.get("estado") or "").lower()]
            total = sum(float(f.get("importe") or 0) for f in validas)
            bg, fg = _COLOR_METODO.get(metodo.lower(), _COLOR_METODO["otro"])
            r += 2
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ult)
            x = ws.cell(row=r, column=1, value=f"{metodo.upper()}   ·   {len(validas)} pago(s)   ·   ${total:,.2f}")
            x.font, x.fill, x.alignment = _font(bold=True, size=12, color=fg), _f(bg), _alin(vertical="center", indent=1)
            ws.row_dimensions[r].height = 26
            r += 1
            encabezado(r, cols_t)
            r = filas_(r, cols_t, filas)
            r += 1
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(i_imp - 1, 1))
            x = ws.cell(row=r, column=1, value=f"Subtotal {metodo}")
            x.font, x.fill, x.alignment = _font(bold=True, size=10), _f(AZUL_CLARO), _alin(horizontal="right")
            y = ws.cell(row=r, column=i_imp, value=total)
            y.number_format, y.font, y.fill, y.border = '"$"#,##0.00', _font(bold=True, size=10, color=AZUL), _f(AZUL_CLARO), BORDE
            ws.row_dimensions[r].height = 22
            resumen_metodos.append((metodo, len(validas), total, len(filas) - len(validas)))
    else:
        encabezado(cab, cols)
        r = filas_(cab, cols, rep["filas"])
        if not rep["filas"]:
            ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=ult)
            x = ws.cell(row=r + 1, column=1, value="No hay información para mostrar con estos filtros.")
            x.font, x.alignment = _font(italic=True, color="667085"), _alin(horizontal="center")
            r += 1
        else:
            ws.auto_filter.ref = f"A{cab}:{get_column_letter(n)}{r}"

    # ---- resumen
    resumen = rep.get("resumen") or []
    if resumen_metodos:
        r += 2
        ws.cell(row=r, column=1, value="RESUMEN POR MÉTODO DE PAGO").font = _font(bold=True, size=10, color=AZUL)
        r += 1
        for j, t in enumerate(("Método", "Pagos", "Total"), 1):
            x = ws.cell(row=r, column=j, value=t)
            x.font, x.fill, x.border, x.alignment = _font(bold=True, color="FFFFFF", size=10), _f(AZUL), BORDE, _alin(horizontal="center")
        for metodo, cant, total, canc in resumen_metodos:
            r += 1
            bg, fg = _COLOR_METODO.get(metodo.lower(), _COLOR_METODO["otro"])
            a = ws.cell(row=r, column=1, value=metodo + (f"  ({canc} cancelado(s) no suman)" if canc else ""))
            a.font, a.fill, a.border = _font(bold=True, size=10, color=fg), _f(bg), BORDE
            ws.cell(row=r, column=2, value=cant).border = BORDE
            ws.cell(row=r, column=2).alignment = _alin(horizontal="center")
            t = ws.cell(row=r, column=3, value=total)
            t.number_format, t.border = '"$"#,##0.00', BORDE
        r += 1
        ws.cell(row=r, column=1, value="TOTAL GENERAL").font = _font(bold=True, size=11, color="FFFFFF")
        for j in (1, 2, 3):
            ws.cell(row=r, column=j).fill, ws.cell(row=r, column=j).border = _f(AZUL), BORDE
        ws.cell(row=r, column=2, value=sum(m[1] for m in resumen_metodos)).font = _font(bold=True, color="FFFFFF")
        ws.cell(row=r, column=2).alignment = _alin(horizontal="center")
        g = ws.cell(row=r, column=3, value=sum(m[2] for m in resumen_metodos))
        g.number_format, g.font = '"$"#,##0.00', _font(bold=True, size=11, color="FFFFFF")
        resumen = []
    if resumen:
        r += 2
        ws.cell(row=r, column=1, value="RESUMEN").font = _font(bold=True, size=10, color=AZUL)
        for k, v in resumen:
            r += 1
            if n >= 3:
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            a = ws.cell(row=r, column=1, value=k)
            a.font, a.fill, a.border = _font(bold=True, size=10), _f(AZUL_CLARO), BORDE
            if n >= 3:
                ws.cell(row=r, column=2).border = BORDE
            b = ws.cell(row=r, column=3 if n >= 3 else 2, value=v)
            b.font, b.border, b.alignment = _font(bold=True, size=10, color=AZUL), BORDE, _alin(horizontal="right")
            ws.row_dimensions[r].height = 20

    # ---- anchos, impresión
    cols_w = cols_t if por_metodo else cols
    for j, c in enumerate(cols_w, 1):
        ancho = max([len(str(c["label"]))] + [len(str(f.get(c["key"]) if f.get(c["key"]) is not None else ""))
                                              for f in rep["filas"][:300]])
        ws.column_dimensions[get_column_letter(j)].width = min(max(ancho + 3, 11), 48)
    if con_logo:
        ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 18)
    if not resumen_metodos:
        ws.freeze_panes = ws.cell(row=cab + 1, column=1)
    else:
        ws.freeze_panes = ws.cell(row=cab, column=1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if not resumen_metodos:
        ws.print_title_rows = f"{cab}:{cab}"
    ws.oddFooter.center.text = "Página &P de &N"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
